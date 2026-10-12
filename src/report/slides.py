"""Build outputs/presentation.pptx from the current pipeline outputs.

All numbers come from the same JSON / raster files the situation report uses, so the
deck cannot drift from the report. Run from anywhere:

    python src/report/slides.py
"""
import datetime
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from src import config  # noqa: E402


def J(rel):
    return json.loads((ROOT / rel).read_text(encoding="utf-8"))


def mask_stats():
    import rasterio
    with rasterio.open(ROOT / "data/processed/flood_mask.tif") as s:
        px = int((s.read(1) > 0).sum())
        area = px * abs(s.res[0] * s.res[1]) / 1e6
    return px, area


imp = J("outputs/impact_summary.json")
cut = J("outputs/cutoff_settlements.json")["summary"]
_vp = ROOT / "outputs/validation/emsr927_validation.json"
val = json.loads(_vp.read_text(encoding="utf-8")) if _vp.exists() else None
px, area = mask_stats()


def ems_lines(val):
    """Validation slide text from emsr927_validation.json (assessed domain, flood_or_debris)."""
    if not val:
        return ["• EMSR927 check not run (data/external/EMSR927 missing)"]
    out = ["• Check-only reference: EMSR927 landslide polygons and damage points (GRA products)"]
    a = val["aois"].get("AOI03", {})
    seg = a.get("event_area", {}).get("flood_or_debris", {}).get("assessed", {}).get("strict")
    if seg and seg.get("precision") is not None:
        out.append(f"• Hold-out AOI03 vs landslide polygons: precision {seg['precision']:.2f}, "
                   f"recall {seg['recall']:.2f}, IoU {seg['iou']:.2f}")
    r0 = (a.get("damage_points", {}).get("recall_assessed", {}) or {}).get("0")
    if r0:
        lo, hi = r0["ci95"]
        out.append(f"• Damage-point recall at 0 m: {r0['rate']:.0%} [{lo:.0%}-{hi:.0%}], {r0['verdict']}")
    out.append("• No flood-extent layer: the flood class is not validated against EMS")
    return out

prs = Presentation()


def slide(title, lines, size=22):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    for p in tf.paragraphs:
        p.font.size = Pt(size)


det = J(f"outputs/diagnostics/detect_{config.TRACK}.json")
_cn = ROOT / "outputs/validation/class_nulls.json"
cn = json.loads(_cn.read_text(encoding="utf-8")) if _cn.exists() else None
_co = ROOT / "outputs/validation/ai_corridor.json"
co = json.loads(_co.read_text(encoding="utf-8")) if _co.exists() else None
b, rr = imp["buildings"], (imp.get("sensitivity") or {}).get("range") or {}
rb = imp["roads_bridges"]["exposure_by_tolerance"]["0m"]
cutj = J("outputs/cutoff_settlements.json"); rc = (cutj.get("sensitivity") or {}).get("range") or {}
def rng(r, k, fmt="{:,}"):
    v = r.get(k)
    return "" if not v else f" (range {fmt.format(v[0])}-{fmt.format(v[1])})"
unk = (b["not_assessed"] + b["unobservable"]) / max(b["total"], 1)
A_ = det["area_km2"]

slide(config.TITLE + ": Sentinel-1 analysis",
      ["Calibrated SAR change detection and OpenStreetMap exposure",
       "Educational prototype, not an operational tool",
       datetime.date.today().isoformat()])

slide("Problem",
      [f"• Event date {config.FLOOD_DATE}, {config.AREA_NAME}",
       "• Can free Sentinel-1 data show where the valley floor changed, and what is exposed?",
       "• Checking only: Copernicus EMS products (EMSR927 GRA: landslide polygons and damage points, no flood-extent layer)"])

slide("Data",
      [f"• Sentinel-1 GRD VV/VH, track {det['track']}: references {', '.join(config.REF_DATES[r] for r in det['refs'])}; post {config.POST_DATE}",
       "• Copernicus GLO-30 DEM (HAND valley floor, radar layover / shadow mask)",
       f"• OpenStreetMap snapshot of {config.OSM_SNAPSHOT_DATE} (pre-event)",
       "• EMS products used only for checking"])

slide("Method",
      ["1. SNAP: orbit file, noise removal, sigma0, terrain correction; 30 m grid",
       f"2. Robust-z change of the post scene against {len(det['refs'])} pre-event references, valley floor only",
       "3. Thresholds from leave-one-out null runs on pre-event dates (no EMS)",
       "4. Overlay on pre-event OSM: exposure, bridges exposed, hospital access"])

slide("Key finding",
      [f"• Detector status: {det['status']} (calibrated = {det['calibrated']})",
       f"• Flagged: {A_['flood_water']:.2f} km² flood-like (water-like backscatter change), {A_['surface_dark']:.2f} km² dark, {A_['surface_bright']:.2f} km² bright",
       "• " + (cn["statement"] if cn else "Class-level nulls not run"),
       "• Surface change on the valley floor, not a validated flood extent"], size=16)

slide("Exposure to the mask (not damage)",
      [f"• Buildings overlapping the mask: {b['overlap']:,}{rng(rr, 'buildings_overlap')} of {b['total']:,}; {unk:.0%} unknown status",
       f"• Motorable road in mask: {rb['motorable_road_km']:.1f} km{rng(rr, 'motorable_road_km', '{:.1f}')}",
       f"• Bridges exposed (not failed): {rb['bridges']}{rng(rr, 'bridges_exposed')}",
       f"• Settlements with no road path to a hospital: {cut.get('cut_off', 0)}{rng(rc, 'hospital_cut_off')} of {cut['total_settlements']}"], size=18)

slide("Validation",
      ems_lines(val) +
      ([f"• Water model on the corridor (prediction-only): {co['verdict']}, so it is not independent evidence"] if co else []), size=18)

slide("Deliverables",
      ["• Interactive dashboard (outputs/dashboard/index.html, works offline)",
       "• Situation reports (EN, NE, one page) and docs/REPORT.pdf",
       "• Reproducible pipeline: python -m src.run (SNAP route; --s1-source pc: cross-check only)",
       "• Water-model weights (models/)"])

slide("Limitations and next steps",
      [f"• One track; post scene {config.POST_DATE}",
       "• 30 m pixels, steep terrain masked, OSM incomplete in rural areas",
       "• Next: a second track or a closer post scene, post-event OSM for cut-off checks"])

out = ROOT / "outputs/presentation.pptx"
prs.save(out)
print("wrote", out)